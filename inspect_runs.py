import docx

doc_path = r"C:\Users\USER\Desktop\doc\interim-layout.docx"
doc = docx.Document(doc_path)

with open("docx_full_paragraphs.txt", "w", encoding="utf-8") as f:
    f.write(f"Total paragraphs: {len(doc.paragraphs)}\n")
    for i, p in enumerate(doc.paragraphs):
        f.write(f"P{i} (Style: {p.style.name}): '{p.text}'\n")
        for r_idx, run in enumerate(p.runs):
            f.write(f"  Run {r_idx}: '{run.text}' (bold: {run.bold}, italic: {run.italic})\n")
    f.write(f"\nTotal tables: {len(doc.tables)}\n")
    for t_idx, t in enumerate(doc.tables):
        f.write(f"Table {t_idx}: {len(t.rows)} rows, {len(t.columns)} cols\n")
